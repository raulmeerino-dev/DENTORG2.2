"""Bounded-memory export writers for the authorized records projection."""
import csv
import io
from collections.abc import Mapping, Sequence
from contextlib import suppress
from datetime import date, datetime
from decimal import Decimal
from tempfile import SpooledTemporaryFile
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import Paragraph

from app.domains.reporting.schemas.registros import RegistroColumn

MIME_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}
MAX_PDF_ROWS = 10_000


class ExportValueError(ValueError):
    """The format cannot represent a value without silently changing it."""


def text_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Sí" if value else "No"
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def csv_value(value):
    """Treat formula-looking source strings as text when opened in a spreadsheet."""
    text = text_value(value)
    if isinstance(value, str) and text.lstrip().startswith(("=", "+", "-", "@", "\t", "\r", "\n")):
        return "'" + text
    return text


class RecordExportWriter:
    def __init__(self, format: str, columns: Sequence[RegistroColumn], title: str, subtitle: str):
        self.format = format
        self.columns = columns
        self.title = title
        self.subtitle = subtitle
        self.count = 0
        self.finished = False
        # Ownership transfers to StreamingResponse; its iterator/background task closes it.
        self.file = SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode="w+b")  # noqa: SIM115
        self.sheet = None
        if format == "csv":
            self.file.write(b"\xef\xbb\xbf")
            self._csv_row([column.label for column in columns])
        elif format == "xlsx":
            self.workbook = Workbook(write_only=True)
            self.workbook.properties.title = title
            self.workbook.properties.description = subtitle
            self._new_sheet()
        elif format == "pdf":
            self.canvas = Canvas(self.file, pagesize=landscape(A4))
            self.canvas.setTitle(title)
            self.canvas.setAuthor("DentCore")
            self.page = 0
            self.width, self.height = landscape(A4)
            weights = [1.7 if column.type == "text" else 1 for column in columns]
            self.widths = [(self.width - 56) * weight / sum(weights) for weight in weights]
            self.style = ParagraphStyle("record", fontName="Helvetica", fontSize=7, leading=9, alignment=TA_LEFT, splitLongWords=True)
            self._new_page()
        else:
            self.file.close()
            raise ValueError("Formato de exportación no admitido")

    def _csv_row(self, values):
        buffer = io.StringIO(newline="")
        csv.writer(buffer, delimiter=";", lineterminator="\r\n").writerow(values)
        self.file.write(buffer.getvalue().encode("utf-8"))

    def _new_sheet(self):
        self.sheet = self.workbook.create_sheet(f"Registros {len(self.workbook.worksheets) + 1}")
        self.sheet.freeze_panes = "A2"
        for index, column in enumerate(self.columns, 1):
            self.sheet.column_dimensions[get_column_letter(index)].width = 32 if column.type == "text" else 21
        header = []
        for column in self.columns:
            cell = WriteOnlyCell(self.sheet, column.label)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="234F64")
            header.append(cell)
        self.sheet.append(header)
        self.sheet_rows = 1

    def _finish_sheet(self):
        self.sheet.auto_filter.ref = f"A1:{get_column_letter(len(self.columns))}{self.sheet_rows}"

    def _new_page(self):
        if self.page:
            self.canvas.showPage()
        self.page += 1
        self.canvas.setFillColor(colors.HexColor("#153642"))
        self.canvas.setFont("Helvetica-Bold", 12)
        self.canvas.drawString(28, self.height - 29, f"DentCore · {self.title}")
        self.canvas.setFont("Helvetica", 8)
        self.canvas.drawString(28, self.height - 43, self.subtitle)
        self.canvas.drawRightString(self.width - 28, 17, f"Página {self.page}")
        self.canvas.setFillColor(colors.HexColor("#e7eff2"))
        self.canvas.rect(28, self.height - 69, self.width - 56, 18, fill=1, stroke=0)
        self.canvas.setFillColor(colors.HexColor("#153642"))
        x = 28
        for column, width in zip(self.columns, self.widths, strict=True):
            paragraph = Paragraph(f"<b>{escape(column.label)}</b>", self.style)
            _, height = paragraph.wrap(width - 8, 30)
            paragraph.drawOn(self.canvas, x + 4, self.height - 56 - height)
            x += width
        self.y = self.height - 74

    def add(self, cells: Mapping):
        values = [cells.get(column.key) for column in self.columns]
        if self.format == "csv":
            self._csv_row([csv_value(value) for value in values])
        elif self.format == "xlsx":
            if self.sheet_rows >= 1_048_576:
                self._finish_sheet()
                self._new_sheet()
            row = []
            for column, value in zip(self.columns, values, strict=True):
                numeric = column.type in {"number", "money"} and isinstance(value, (int, float, Decimal)) and not isinstance(value, bool)
                text = text_value(value)
                if not numeric and (len(text) > 32_767 or ILLEGAL_CHARACTERS_RE.search(text)):
                    raise ExportValueError("Una celda supera los límites de texto de Excel. Use CSV para conservar su contenido completo.")
                cell = WriteOnlyCell(self.sheet, value if numeric else text)
                if not numeric:
                    cell.data_type = "s"  # Never interpret patient-entered strings as formulas.
                elif column.type == "money":
                    cell.number_format = '#,##0.00 "€"'
                row.append(cell)
            self.sheet.append(row)
            self.sheet_rows += 1
        else:
            paragraphs = [Paragraph(escape(text_value(value)).replace("\n", "<br/>"), self.style) for value in values]
            # Split long cells across pages instead of clipping clinical/document descriptions.
            while any(paragraphs):
                if self.y - 32 < 24:
                    self._new_page()
                available = self.y - 32
                pieces, tails, heights = [], [], []
                for paragraph, width in zip(paragraphs, self.widths, strict=True):
                    if paragraph is None:
                        pieces.append(None)
                        tails.append(None)
                        heights.append(0)
                        continue
                    _, height = paragraph.wrap(width - 8, available)
                    split = paragraph.split(width - 8, available - 8) if height > available - 8 else [paragraph]
                    piece = split[0] if split else None
                    pieces.append(piece)
                    tails.append(split[1] if len(split) > 1 else paragraph if not split else None)
                    heights.append(piece.wrap(width - 8, available)[1] if piece else 0)
                row_height = max(heights, default=0) + 8
                x = 28
                for piece, height, width in zip(pieces, heights, self.widths, strict=True):
                    if piece:
                        piece.drawOn(self.canvas, x + 4, self.y - height - 4)
                    x += width
                self.y -= row_height
                self.canvas.setStrokeColor(colors.HexColor("#d7e0e4"))
                self.canvas.line(28, self.y, self.width - 28, self.y)
                paragraphs = tails
                if any(paragraphs):
                    self._new_page()
        self.count += 1

    def finish(self):
        if self.format == "xlsx":
            self._finish_sheet()
            self.workbook.save(self.file)
            self.workbook.close()
        elif self.format == "pdf":
            self.canvas.save()
        self.file.seek(0)
        self.finished = True
        return self.file

    def close(self):
        if self.format == "xlsx" and not self.finished:
            # Finalize the write-only sheet's temporary XML even on a cancelled export.
            with suppress(Exception):
                self.finish()
        self.file.close()
